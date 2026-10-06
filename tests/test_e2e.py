"""Cadena lógica completa (Plata -> MDM -> Oro -> cruces -> liquidación -> listados -> informe) sobre
datos sintéticos, sin Spark. Verifica los criterios de la PoC antes de llevarla a Fabric."""
import sys
from datetime import date
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "img_lib" / "src"))
sys.path.insert(0, str(ROOT / "generator"))

from img_lib.validate import validate_record

import generator
from img_lib import active, crosschecks, dispersal, mapping, mdm, metrics, settlement
from img_lib import illustrative as ill

TODAY = date(2026, 10, 1)
CUTOFF = "2026-09-30"


def run_chain(n=500, cutoff=1):
    data = generator.generate(n, cutoff)
    silver, quarantine = [], []
    for source in ("population", "validation"):
        for r in data[source]:
            rec, causes = validate_record({**mapping.to_internal(r, source), "source": source}, today=TODAY)
            (quarantine if causes else silver).append(rec)
    people, xref = mdm.build_master(silver)
    _, assign = mdm.build_households(people)
    person_of = {(x["source"], x["origin_id"]): x["person_id"] for x in xref}
    gold = {"population": [], "validation": []}
    for r in silver:
        pid = person_of[(r["source"], r["origin_id"])]
        gold[r["source"]].append({**r, "person_id": pid, "household_id": assign[pid]})
    base = crosschecks.build_base(gold["population"], gold["validation"], cycle_id="2026-09", cutoff_date=CUTOFF,
                                 source_versions={"population": "c1", "validation": "c1"},
                                 execution_id="E-test", notebook_version="test", ts="2026-10-01T00:00:00")
    p = {k: active(getattr(ill, k), CUTOFF) for k in
         ("TARGETING_CRITERIA", "BLOCK_RULES", "HOLDER_RULE", "OPERATORS", "AMOUNTS", "FUNDING_SOURCES")}
    tgt = settlement.target(base, p["TARGETING_CRITERIA"], p["BLOCK_RULES"])
    hld = settlement.select_holder(tgt, base, p["HOLDER_RULE"])
    mp = settlement.assign_payment_method(hld, p["OPERATORS"], base)
    amt = settlement.calculate_amount(hld, p["AMOUNTS"])
    payments, status = settlement.assign_funding_source(amt, p["FUNDING_SOURCES"])
    master_sheet, files = settlement.generate_payment_lists(payments, mp, base)
    return dict(data=data, mp=mp, quarantine=quarantine, xref=xref, base=base, p=p, tgt=tgt, hld=hld, payments=payments,
                status=status, master_sheet=master_sheet, files=files)


def test_full_chain_and_criteria():
    c = run_chain()
    v = c["data"]["ground_truth"]
    # base de cruces: una fila por persona, con los grupos de la Tabla 7
    ids = [r["person_id"] for r in c["base"]]
    assert len(ids) == len(set(ids))
    for field in ("cycle_id", "person_id", "household_id", "sisben_group", "validation_found", "execution_id",
                  "sex", "age", "banked", "household_role", "renec_validity", "operator"):
        assert field in c["base"][0]
    # un titular por hogar con titular elegible (adulta/o, no bloqueada/o); los demás hogares quedan sin titular
    by_person = {r["person_id"]: r for r in c["base"]}
    eligible = {f["person_id"] for f in c["tgt"] if f["eligible"]}
    assert len({t["household_id"] for t in c["hld"]}) == len(c["hld"]) > 0
    assert all(t["person_id"] in eligible and by_person[t["person_id"]]["age"] >= 18 for t in c["hld"])
    eligible_adult_households = {by_person[p]["household_id"] for p in eligible if by_person[p]["age"] >= 18}
    assert {t["household_id"] for t in c["hld"]} == eligible_adult_households
    # focalización por grupo A; bloqueos por RENEC (NOT IN 0,12) e inhumados
    assert all(by_person[p]["sisben_group"] == "1. SISBEN IV - A" for p in eligible)
    reasons = {x for f in c["tgt"] for x in (f["reasons"] or "").split(";") if x}
    assert {"RENEC_NOT_VALID", "IN_DECEASED_REGISTRY"} <= reasons
    # error de bloqueo y de titular < 1 % contra la verdad conocida
    def expected(i):
        s = set()
        if i["renec_blocked"]:
            s.add("RENEC_NOT_VALID")
        if i["in_registry"]:
            s.add("IN_DECEASED_REGISTRY")
        return s
    block_reasons = [b["reason"] for b in c["p"]["BLOCK_RULES"]]
    assert metrics.block_error_rate(v, c["xref"], c["tgt"], expected, block_reasons) < 0.01
    assert metrics.holder_error_rate(v, c["xref"], c["hld"]) < 0.01
    # suma de listados == liquidación; el monto es el del diccionario de dispersión
    assert abs(sum(f["amount"] for f in c["master_sheet"]) - sum(p["amount"] for p in c["payments"])) < 1e-6
    assert {p["amount"] for p in c["payments"]} == {120000.0}
    assert 0 < crosschecks.match_pct(c["base"]) < 100


def test_origin_id_reaches_the_dispersal_listing():
    """sdp_id_llave_maestra sale de origin_id: la base de cruces debe conservarlo hasta el listado."""
    c = run_chain()
    assert all(r["origin_id"] for r in c["base"])
    assert all(f["origin_id"] for f in c["master_sheet"])
    llaves = [dispersal.to_dispersal_row(f)["sdp_id_llave_maestra"] for f in c["master_sheet"]]
    assert llaves and all(isinstance(x, int) for x in llaves)


def test_report_publication_and_fingerprint(tmp_path):
    c = run_chain()
    inf = settlement.build_report(
        cycle_id="2026-09", cutoff_date=CUTOFF, source_versions={"population": "c1"},
        applied_params=c["p"], notebook_versions={"nb_00": "test"}, base=c["base"],
        targeting=c["tgt"], holders=c["hld"], payments=c["payments"], master_sheet=c["master_sheet"], approvals=[])
    assert inf["uses_illustrative_params"] is True  # R8: debe verse en el informe
    assert inf["lists_sum_equals_settlement"] is True
    assert inf["counts"]["settled_payments"] == len(c["hld"])
    man = settlement.publish(c["files"], str(tmp_path), "2026-09")
    assert man and (tmp_path / "2026-09" / "manifest_sha256.csv").exists()
    from openpyxl import load_workbook

    from img_lib import sha256_file
    assert {m["file"].split("/")[0] for m in man} == set(c["files"])  # un archivo .xlsx por operador
    total = 0
    for m in man:
        assert m["file"].endswith(".xlsx")
        assert sha256_file(str(tmp_path / "2026-09" / m["file"])) == m["sha256"]
        ws = load_workbook(tmp_path / "2026-09" / m["file"]).active
        rows = list(ws.iter_rows(values_only=True))
        assert list(rows[0]) == dispersal.COLUMNS and len(rows[0]) == 32
        assert len(rows) - 1 == m["rows"]
        operator = m["file"].split("/")[0]
        assert {r[dispersal.COLUMNS.index("sdp_operador")] for r in rows[1:]} == {operator}
        total += m["rows"]
    assert total == len(c["master_sheet"])


def test_audit_coverage():
    assert metrics.audit_coverage(["E1", "E1", "E2"], ["E1", "E2"]) == 1.0
    assert metrics.audit_coverage(["E1", "E3"], ["E1"]) == 0.5


def test_mdm_and_quarantine_error_at_5000():
    c = run_chain(5000)
    v = c["data"]["ground_truth"]
    assert metrics.mdm_error_rate(v, c["xref"]) < 0.01
    ids = {(r["source"], r["origin_id"]) for r in c["quarantine"]}
    assert metrics.quarantine_error_rate(v, ids) < 0.01
