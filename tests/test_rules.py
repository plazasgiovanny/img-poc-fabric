"""Reglas de focalización, bloqueo y titular, y estructura del listado de dispersión."""
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "img_lib" / "src"))

from img_lib.validate import validate_record
from openpyxl import load_workbook

from img_lib import active, crosschecks, dispersal, mapping, mdm, settlement
from img_lib import illustrative as ill

CUTOFF = "2026-09-30"
HOLDER_RULE = active(ill.HOLDER_RULE, CUTOFF)


def person(pid, household="H1", age=30, sex="1", banked=0, doc="99000001", renec="0", group=mapping.GROUP_TARGET,
           **kw):
    return {"cycle_id": "2026-09", "person_id": pid, "household_id": household, "age": age, "sex": sex,
            "banked": banked, "doc_number": doc, "doc_type": "1", "renec_validity": renec, "sisben_group": group,
            "validation_found": "NO", "operator": "SIN OPERADOR", "origin_id": "7" + pid[-6:].zfill(7), **kw}


def holders(members):
    tgt = settlement.target(members, active(ill.TARGETING_CRITERIA, CUTOFF), active(ill.BLOCK_RULES, CUTOFF))
    return {h["household_id"]: h["person_id"] for h in settlement.select_holder(tgt, members, HOLDER_RULE)}


def reasons(p):
    return settlement.target([p], active(ill.TARGETING_CRITERIA, CUTOFF), active(ill.BLOCK_RULES, CUTOFF))[0]["reasons"]


def test_targeting_by_group_equals():
    assert reasons(person("P1")) is None
    assert reasons(person("P2", group="2. SISBEN IV - B")) == "NOT_MET_SISBEN_GROUP_A"


def test_block_renec_not_in_operator():
    for code in ("0", "12"):
        assert reasons(person("P1", renec=code)) is None  # vigente: no bloquea
    for code in ("21", "22", "99"):
        assert reasons(person("P1", renec=code)) == "RENEC_NOT_VALID"  # cualquier otro código bloquea
    assert reasons(person("P1", renec=None)) is None  # nula = no bloqueada
    row = {"x": "5"}
    assert settlement._matches(row, "x", "NOT IN", "0,12") and not settlement._matches(row, "x", "NOT IN", "5, 6")
    assert not settlement._matches({"x": None}, "x", "NOT IN", "0,12")


def _base_with_deceased(registry_type):
    recs = []
    for source, doc_type in (("population", "1"), ("validation", registry_type)):
        if source == "population":
            raw = {"RSH_id_llave_maestra": "70000001", "RSH_tip_documento": doc_type, "RSH_num_documento": "99111111",
                   "RSH_pri_nombre": "ANA", "RSH_pri_apellido": "RUIZ", "RSH_sexo_persona": "2", "SIS_edad": "40",
                   "RSH_fec_nacimiento": "1986-01-01", "RSH_grupo_S4": mapping.GROUP_TARGET, "SIS_cod_loc": "1", "RSH_id_hogar": "H1"}
        else:
            raw = {"RSH_tip_documento": doc_type, "RSH_num_documento": "99111111", "fecha_defuncion": "2026-01-01"}
        rec, causes = validate_record({**mapping.to_internal(raw, source), "source": source})
        assert causes == []
        recs.append(rec)
    people, xref = mdm.build_master(recs)
    _, assign = mdm.build_households(people)
    pid = {(x["source"], x["origin_id"]): x["person_id"] for x in xref}
    gold = {s: [{**r, "person_id": pid[(r["source"], r["origin_id"])],
                 "household_id": assign[pid[(r["source"], r["origin_id"])]]}
                for r in recs if r["source"] == s] for s in ("population", "validation")}
    return crosschecks.build_base(gold["population"], gold["validation"], cycle_id="2026-09", cutoff_date=CUTOFF,
                                  source_versions={}, execution_id="E", notebook_version="t", ts="t")


def test_deceased_registry_matches_on_type_and_number():
    same = _base_with_deceased("1")
    assert same[0]["validation_found"] == "YES" and same[0]["death_date"] == "2026-01-01"
    assert "IN_DECEASED_REGISTRY" in reasons(same[0])
    other_type = _base_with_deceased("3")  # mismo número, otro tipo: otra identidad
    assert other_type[0]["validation_found"] == "NO"
    assert reasons(other_type[0]) is None


def test_holder_must_be_adult_and_not_blocked():
    m = [person("P1", age=17, sex="2"), person("P2", age=40, sex="1", renec="21"), person("P3", age=18, sex="1")]
    assert holders(m) == {"H1": "P3"}  # menor y bloqueado descartados
    assert holders([person("P1", age=17), person("P2", age=10)]) == {}  # hogar sin titular elegible
    assert holders([person("P1", age=50, renec="21")]) == {}


def test_holder_prefers_woman_then_banked_then_older_then_document():
    man_banked = person("P1", sex="1", banked=1, age=60, doc="99000001")
    woman = person("P2", sex="2", banked=0, age=25, doc="99000002")
    assert holders([man_banked, woman]) == {"H1": "P2"}  # mujer antes que bancarizado
    woman_banked = person("P3", sex="2", banked=1, age=22, doc="99000003")
    assert holders([man_banked, woman, woman_banked]) == {"H1": "P3"}  # bancarizada: preferencia, no descarte
    assert holders([woman, person("P4", sex="2", banked=0, age=70, doc="99000009")]) == {"H1": "P4"}  # mayor edad
    older_tie = [person("P5", sex="2", age=40, doc="99000050"), person("P6", sex="2", age=40, doc="99000010")]
    assert holders(older_tie) == {"H1": "P6"}  # documento ascendente
    assert holders([woman]) == {"H1": "P2"}  # no bancarizada también puede ser titular
    men = [person("P7", sex="1", banked=0, age=70), person("P8", sex="1", banked=1, age=30, doc="99000008")]
    assert holders(men) == {"H1": "P8"}  # sin mujer adulta: bancarizado antes que mayor


def test_holder_ignores_non_focalized_members_and_splits_households():
    m = [person("P1", group="2. SISBEN IV - B", sex="2"), person("P2", sex="1"),
         person("P3", household="H2", sex="2")]
    assert holders(m) == {"H1": "P2", "H2": "P3"}


def test_payment_method_account_operator_with_fallback():
    base = [person("P1", operator="NEQUI"), person("P2", operator="SIN OPERADOR")]
    hld = [{"cycle_id": "2026-09", "household_id": "H1", "person_id": "P1"},
           {"cycle_id": "2026-09", "household_id": "H2", "person_id": "P2"}]
    mp = settlement.assign_payment_method(hld, active(ill.OPERATORS, CUTOFF), base)
    assert [(m["operator"], m["mode"]) for m in mp] == [("NEQUI", "account_operator"), ("EFECTY", "fallback_operator")]


def _sheet_row(**over):
    row = {"cycle_id": "2026-09", "seq": 7, "operator": "DAVIVIENDA", "product": 1, "origin_id": "70000042",
           "doc_type": "1", "doc_number": "99123456", "first_name": "ANA", "second_name": None,
           "last_name": "RUIZ", "second_last_name": "PEÑA", "renec_validity": "0", "age": 40, "sex": "2",
           "locality": "999", "locality_name": "ZZ-SIN INFORMACION", "amount": 120000.0}
    row.update(over)
    return row


def test_dispersal_row_has_32_columns_in_order_and_defaults():
    assert len(dispersal.COLUMNS) == 32 and len(set(dispersal.COLUMNS)) == 32
    assert dispersal.COLUMNS[0] == "sdp_id_listado" and dispersal.COLUMNS[-1] == "parqueadero"
    r = dispersal.to_dispersal_row(_sheet_row())
    assert list(r) == dispersal.COLUMNS
    assert r["sdp_tip_documento"] == 3 and r["sdp_tip_documento_texto"] == "Cedula de Ciudadania"  # CC maestra 1 -> 3
    assert r["sdp_id_listado"] == 1202609000007 and r["sdp_id_pago"] == "IMG-202609-000007"
    assert r["sdp_id_IMG"] == 2202609000007 and r["sdp_Giro"] == 20260901
    assert r["sdp_cod_localidad"] == 999 and r["sdp_localidad"] == "ZZ- SIN INFORMACION"
    assert r["sdp_cod_upz"] == 999 and r["sdp_upz"] == "ZZ- SIN INFORMACION"
    assert r["sdp_ncuenta"] == 0 and r["sdp_cel_beneficiario_origen"] == 0 and r["parqueadero"] == "SIN INFORMACION"
    assert r["sdp_seg_nombre"] is None and r["sdp_pri_apellido"] == "RUIZ" and r["sdp_monto"] == 120000.0
    # sin datos: valores por defecto del diccionario
    e = dispersal.to_dispersal_row(_sheet_row(doc_type=None, doc_number=None, locality=None, locality_name=None))
    assert e["sdp_tip_documento"] == 0 and e["sdp_tip_documento_texto"] == "SIN INFORMACION"
    assert e["sdp_num_documento"] == 0 and e["sdp_cod_localidad"] == 999
    assert e["sdp_localidad"] == "ZZ- SIN INFORMACION"
    assert dispersal.to_dispersal_row(_sheet_row(doc_type="4"))["sdp_tip_documento"] == 1  # RC


def test_dispersal_ids_are_deterministic():
    assert dispersal.listing_ids("2026-09", 1) == dispersal.listing_ids("2026-09", 1)
    assert dispersal.listing_ids("2026-10", 1)["sdp_id_pago"] != dispersal.listing_ids("2026-09", 1)["sdp_id_pago"]


def test_xlsx_structure(tmp_path):
    rows_in = [dispersal.to_dispersal_row(_sheet_row()), dispersal.to_dispersal_row(_sheet_row(seq=8))]
    f = tmp_path / "x.xlsx"
    f.write_bytes(dispersal.xlsx_bytes(rows_in))
    rows = list(load_workbook(f).active.iter_rows(values_only=True))
    assert list(rows[0]) == dispersal.COLUMNS and len(rows) == 3


def test_ddl_param_script_matches_python_mirror():
    sql = (ROOT / "ddl" / "02_param_illustrative.sql").read_text(encoding="utf-8")
    for row in ill.HOLDER_RULE:
        assert f"'{row['criterion']}', '{row['sort_direction']}'" in sql
    for row in ill.BLOCK_RULES:
        assert f"'{row['rule_id']}'" in sql and f"'{row['reason']}'" in sql
    assert "'NOT IN', '0,12'" in sql and "120000.0" in sql and "'1. SISBEN IV - A'" in sql
